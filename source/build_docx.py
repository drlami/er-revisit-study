# -*- coding: utf-8 -*-
"""
Build a .docx from one of this study's markdown documents, in the house style
used by the other designed documents in this folder.

    python build_docx.py TEAM-BRIEFING.md "Team Briefing.docx"

House style (matched from Protocol _ Proposal .docx):
    page      Letter, 0.8in side margins, 0.7in top/bottom
    body      Aptos 10.5pt
    H1 16pt bold #17365D   H2 12.5pt bold #2F5597   H3 11pt bold #5B9BD5

Handles: headings, **bold**, *italic*, `code`, bullet lists, numbered lists,
pipe tables, > blockquotes, and --- rules. Written to be re-run: it always
builds a fresh file, so diff against any hand-edited copy before overwriting.
"""
import re
import sys
import io

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BODY_FONT = "Aptos"
MONO_FONT = "Consolas"
H_COLORS = {1: "17365D", 2: "2F5597", 3: "5B9BD5"}
HEADER_SHADE = "DEE7F0"
QUOTE_SHADE = "F2F2F2"


# ---------------------------------------------------------------- inline text

TOKEN = re.compile(r"(\*\*.+?\*\*|(?<!\*)\*(?!\*).+?(?<!\*)\*(?!\*)|`[^`]+`)", re.S)


def add_inline(paragraph, text, bold=False, italic=False, size=None):
    """Write markdown inline formatting into a paragraph."""
    text = text.replace("\\|", "|")
    for part in TOKEN.split(text):
        if not part:
            continue
        b, i, mono, body = bold, italic, False, part
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            b, body = True, part[2:-2]
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            mono, body = True, part[1:-1]
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            i, body = True, part[1:-1]
        run = paragraph.add_run(body)
        run.bold = b
        run.italic = i
        run.font.name = MONO_FONT if mono else BODY_FONT
        if size is not None:
            run.font.size = size
        if mono:
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)


def shade(cell_or_para, hexcolor):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hexcolor)
    target = cell_or_para._tc.get_or_add_tcPr() if hasattr(cell_or_para, "_tc") \
        else cell_or_para._p.get_or_add_pPr()
    target.append(el)


# ------------------------------------------------------------------ document

def new_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.8)
    sec.top_margin = sec.bottom_margin = Inches(0.7)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.13

    for lvl, color in H_COLORS.items():
        st = doc.styles["Heading %d" % lvl]
        st.font.name = BODY_FONT
        st.font.bold = True
        st.font.size = Pt({1: 16, 2: 12.5, 3: 11}[lvl])
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt({1: 18, 2: 14, 3: 11}[lvl])
        st.paragraph_format.space_after = Pt(5)
        st.paragraph_format.keep_with_next = True
    return doc


def add_table(doc, rows):
    # A markdown table written with an empty header row (| | |) is a two-column
    # label/value block, not a headed table: drop the blank row and render all
    # remaining rows as body.
    headed = any(c.strip() for c in rows[0])
    if not headed:
        rows = rows[1:]
    ncols = max(len(r) for r in rows)
    body = rows[1:] if headed else rows
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    if headed:
        for j, text in enumerate(rows[0]):
            cell = table.rows[0].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, text, bold=True, size=Pt(9.5))
            shade(cell, HEADER_SHADE)
    for i, row in enumerate(body, start=1 if headed else 0):
        for j in range(ncols):
            cell = table.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, row[j] if j < len(row) else "", size=Pt(9.5))
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


def add_meta_table(doc, pairs):
    """Front-matter label/value block, borderless."""
    table = doc.add_table(rows=len(pairs), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, (label, value) in enumerate(pairs):
        for j, (text, bold) in enumerate(((label, True), (value, False))):
            cell = table.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            add_inline(p, text, bold=bold, size=Pt(10))
        table.rows[i].cells[0].width = Inches(2.1)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


def split_row(line):
    line = line.strip().strip("|")
    return [c.strip() for c in re.split(r"(?<!\\)\|", line)]


def is_sep(line):
    return bool(re.match(r"^\s*\|?[\s:|-]+\|[\s:|-]*$", line)) and "-" in line


def convert(md_path, docx_path):
    lines = io.open(md_path, encoding="utf-8").read().split("\n")
    doc = new_document()
    i, n = 0, len(lines)

    while i < n:
        line = lines[i]
        stripped = line.strip()

        if not stripped or stripped in ("---", "***", "___"):
            i += 1
            continue

        # table
        if stripped.startswith("|") and i + 1 < n and is_sep(lines[i + 1]):
            rows = [split_row(stripped)]
            i += 2
            while i < n and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue

        # heading
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = min(len(m.group(1)), 3)
            text = m.group(2).strip()
            p = doc.add_paragraph(style="Heading %d" % level)
            add_inline(p, text)
            i += 1
            continue

        # front-matter block: consecutive "**Label:** value" lines become a
        # borderless two-column table, as in the earlier designed documents
        if re.match(r"^\*\*[^*]+:\*\*", stripped):
            meta = []
            while i < n and re.match(r"^\*\*[^*]+:\*\*", lines[i].strip()):
                m2 = re.match(r"^\*\*([^*]+):\*\*\s*(.*)$", lines[i].strip())
                meta.append((m2.group(1).strip(), m2.group(2).strip()))
                i += 1
            if len(meta) >= 2:
                add_meta_table(doc, meta)
            else:
                p = doc.add_paragraph()
                add_inline(p, "**%s:** %s" % meta[0])
            continue

        # blockquote
        if stripped.startswith(">"):
            buf = []
            while i < n and lines[i].strip().startswith(">"):
                buf.append(re.sub(r"^\s*>\s?", "", lines[i]))
                i += 1
            for chunk in [c for c in "\n".join(buf).split("\n\n") if c.strip()]:
                p = doc.add_paragraph()
                pf = p.paragraph_format
                pf.left_indent = Inches(0.22)
                pf.right_indent = Inches(0.12)
                pf.space_before = Pt(4)
                pf.space_after = Pt(8)
                add_inline(p, chunk.replace("\n", " ").strip())
                shade(p, QUOTE_SHADE)
            continue

        # bullet / numbered list
        m = re.match(r"^(\s*)([-*+]|\d+\.)\s+(.*)$", line)
        if m:
            indent = len(m.group(1))
            ordered = m.group(2)[0].isdigit()
            style = "List Number" if ordered else "List Bullet"
            p = doc.add_paragraph(style=style)
            p.paragraph_format.left_indent = Inches(0.28 + 0.25 * (indent // 2))
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, m.group(3).strip())
            i += 1
            continue

        # paragraph (join soft-wrapped lines)
        buf = [stripped]
        i += 1
        while i < n and lines[i].strip() and not re.match(
                r"^\s*(#{1,6}\s|[-*+]\s|\d+\.\s|\||>|---)", lines[i]):
            buf.append(lines[i].strip())
            i += 1
        p = doc.add_paragraph()
        add_inline(p, " ".join(buf))

    doc.save(docx_path)
    return docx_path


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)
    out = convert(sys.argv[1], sys.argv[2])
    print("built: %s" % out)
